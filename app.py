from __future__ import annotations

import hashlib
import hmac
import json
import logging
import os
import re
import urllib.parse
from contextlib import asynccontextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import requests
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse
from PIL import Image, ImageOps
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook

try:
    from supabase import Client as SupabaseClient
    from supabase import create_client as create_supabase_client
except Exception:
    SupabaseClient = None
    create_supabase_client = None


load_dotenv()
logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO").upper(),
    format="%(asctime)s %(levelname)s %(name)s - %(message)s",
)
logger = logging.getLogger("aadhar_scanner.main")

AZURE_OCR_URL = "https://portal.vision.cognitive.azure.com/api/demo/analyze?features=read"
AZURE_CAPTION_URL = "https://portal.vision.cognitive.azure.com/api/demo/analyze?features=denseCaptions"
AZURE_HEADERS = {"api-call-origin": "Microsoft.Cognitive.CustomVision.Portal"}

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
TELEGRAM_WEBHOOK_SECRET = os.getenv("TELEGRAM_WEBHOOK_SECRET", "").strip()
TELEGRAM_API_BASE = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}" if TELEGRAM_BOT_TOKEN else ""
TELEGRAM_FILE_BASE = f"https://api.telegram.org/file/bot{TELEGRAM_BOT_TOKEN}" if TELEGRAM_BOT_TOKEN else ""

SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_KEY = (
    os.getenv("SUPABASE_SERVICE_KEY", "").strip()
    or os.getenv("SUPABASE_KEY", "").strip()
)
SUPABASE_PARSE_TABLE = os.getenv("SUPABASE_PARSE_TABLE", "aadhaar_parsed").strip()
LOCAL_EXCEL_ENABLED = os.getenv("LOCAL_EXCEL_ENABLED", "false").strip().lower() == "true"
LOCAL_EXCEL_FILE = os.getenv("LOCAL_EXCEL_FILE", "aadhaar_parsed.xlsx").strip()
LOCAL_EXCEL_SHEET = os.getenv("LOCAL_EXCEL_SHEET", "aadhaar_parsed").strip()

MAX_UPLOAD_SIZE_BYTES = int(os.getenv("MAX_UPLOAD_SIZE_BYTES", str(5 * 1024 * 1024)))
ALLOWED_CONTENT_TYPES = {"image/jpeg", "image/png", "image/webp"}
ALLOWED_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}
CORS_ALLOW_ORIGINS = [o.strip() for o in os.getenv("CORS_ALLOW_ORIGINS", "*").split(",") if o.strip()]
ADMIN_TELEGRAM_USER_IDS = {
    int(item.strip())
    for item in os.getenv("ADMIN_TELEGRAM_USER_IDS", "").split(",")
    if item.strip().isdigit()
}

session = requests.Session()


class ParseResult(BaseModel):
    name: str
    dob: str
    uid: str
    gender: str


class ParseListResponse(BaseModel):
    total: int
    is_admin: bool
    records: list[dict[str, Any]]


@dataclass
class TelegramUserCtx:
    user_id: int
    username: str | None
    first_name: str | None
    is_admin: bool


# Verhoeff tables
_v_multiplication = [
    [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
    [1, 2, 3, 4, 0, 6, 7, 8, 9, 5],
    [2, 3, 4, 0, 1, 7, 8, 9, 5, 6],
    [3, 4, 0, 1, 2, 8, 9, 5, 6, 7],
    [4, 0, 1, 2, 3, 9, 5, 6, 7, 8],
    [5, 9, 8, 7, 6, 0, 4, 3, 2, 1],
    [6, 5, 9, 8, 7, 1, 0, 4, 3, 2],
    [7, 6, 5, 9, 8, 2, 1, 0, 4, 3],
    [8, 7, 6, 5, 9, 3, 2, 1, 0, 4],
    [9, 8, 7, 6, 5, 4, 3, 2, 1, 0],
]
_v_permutation = [
    [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
    [1, 5, 7, 6, 2, 8, 3, 0, 9, 4],
    [5, 8, 0, 3, 7, 9, 6, 1, 4, 2],
    [8, 9, 1, 6, 0, 4, 3, 5, 2, 7],
    [9, 4, 5, 3, 1, 2, 6, 8, 7, 0],
    [4, 2, 8, 6, 5, 7, 3, 9, 0, 1],
    [2, 7, 9, 3, 8, 0, 6, 4, 1, 5],
    [7, 0, 4, 6, 9, 1, 3, 2, 5, 8],
]


def validate_verhoeff(number: str) -> bool:
    try:
        if not number.isdigit() or len(number) != 12:
            return False
        if number[0] in {"0", "1"}:
            return False
        c = 0
        for i, item in enumerate(reversed(number)):
            c = _v_multiplication[c][_v_permutation[i % 8][int(item)]]
        return c == 0
    except Exception:
        return False


def extract_valid_uid_from_text(text: str) -> str:
    digit_stream = re.sub(r"\D", "", text or "")
    if len(digit_stream) < 12:
        return ""
    for i in range(len(digit_stream) - 11):
        candidate = digit_stream[i : i + 12]
        if validate_verhoeff(candidate):
            return candidate
    return ""


def _extension(filename: str | None) -> str:
    if not filename or "." not in filename:
        return ""
    return filename.rsplit(".", 1)[1].lower()


def _looks_like_image(data: bytes) -> bool:
    return any(
        (
            data.startswith(b"\xff\xd8\xff"),
            data.startswith(b"\x89PNG\r\n\x1a\n"),
            data.startswith(b"RIFF") and b"WEBP" in data[:16],
        )
    )


def _escape_markdown(text: str) -> str:
    return re.sub(r"([_*`\[])", r"\\\1", text or "")


def telegram_api(method: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    if not TELEGRAM_BOT_TOKEN:
        raise RuntimeError("TELEGRAM_BOT_TOKEN is missing")
    resp = session.post(f"{TELEGRAM_API_BASE}/{method}", json=payload or {}, timeout=60)
    if resp.status_code != 200:
        logger.error("Telegram API Error: %s - %s", resp.status_code, resp.text)
    resp.raise_for_status()
    return resp.json()


def send_message(chat_id: int | str, text: str, reply_to_message_id: int | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "chat_id": str(chat_id),
        "text": text,
        "parse_mode": "Markdown",
    }
    if reply_to_message_id is not None:
        payload["reply_to_message_id"] = reply_to_message_id
    return telegram_api("sendMessage", payload)


def edit_message_text(chat_id: int | str, message_id: int, text: str) -> None:
    telegram_api(
        "editMessageText",
        {
            "chat_id": str(chat_id),
            "message_id": message_id,
            "text": text,
            "parse_mode": "Markdown",
        },
    )


def get_file_bytes(file_id: str) -> bytes:
    meta = telegram_api("getFile", {"file_id": file_id})
    file_path = meta["result"]["file_path"]
    file_url = f"{TELEGRAM_FILE_BASE}/{file_path}"
    response = session.get(file_url, timeout=60)
    response.raise_for_status()
    return response.content


class OCRParser:
    def __init__(self) -> None:
        self.session = requests.Session()

    def detect_card(self, image_bytes: bytes):
        response = self.session.post(
            AZURE_CAPTION_URL,
            headers=AZURE_HEADERS,
            files={"file": ("image.jpg", image_bytes, "image/jpeg")},
            timeout=60,
        )
        response.raise_for_status()
        values = response.json().get("denseCaptionsResult", {}).get("values", [])
        matches = []
        for item in values:
            text = item.get("text", "").lower()
            if any(token in text for token in ("id card", "identity card", "card", "document")):
                box = item.get("boundingBox", {})
                matches.append({"item": item, "area": box.get("w", 0) * box.get("h", 0)})
        if not matches:
            return None
        return min(matches, key=lambda x: x["area"])["item"]["boundingBox"]

    @staticmethod
    def crop_image_bytes(image_bytes: bytes, x: int, y: int, w: int, h: int) -> bytes:
        img = Image.open(BytesIO(image_bytes))
        img = ImageOps.exif_transpose(img)
        img_w, img_h = img.size
        x = max(0, int(x))
        y = max(0, int(y))
        right = min(x + int(w), img_w)
        bottom = min(y + int(h), img_h)
        cropped = img.crop((x, y, right, bottom))
        output = BytesIO()
        cropped.save(output, format="JPEG")
        return output.getvalue()

    def detect_text(self, image_bytes: bytes) -> tuple[dict[str, Any], bytes]:
        cropped = image_bytes
        try:
            box = self.detect_card(image_bytes)
            if box:
                cropped = self.crop_image_bytes(image_bytes, box["x"], box["y"], box["w"], box["h"])
        except Exception as exc:
            logger.warning("Card detection/cropping failed: %s", exc)

        response = self.session.post(
            AZURE_OCR_URL,
            headers=AZURE_HEADERS,
            files={"file": ("image.jpg", cropped, "image/jpeg")},
            timeout=60,
        )
        response.raise_for_status()
        return response.json(), cropped

    @staticmethod
    def extract_full_text(ocr_data: dict[str, Any]) -> str:
        lines: list[str] = []
        for block in ocr_data.get("readResult", {}).get("blocks", []):
            for line in block.get("lines", []):
                text = line.get("text", "")
                if text:
                    lines.append(text)
        return " ".join(lines).strip()

    @staticmethod
    def parse_with_regex(text: str) -> ParseResult:
        name = ""
        dob = ""
        uid = ""
        gender = ""

        name_match = re.search(r"[I|1|l]ndia[^a-zA-Z]*([a-zA-Z\s]+?)[^a-zA-Z]*D[O|0|o]B", text, flags=re.DOTALL)
        if name_match:
            name = name_match.group(1).strip()

        dob_match = re.search(r"D[O|0|o]B[:\s]+(\d{2}\/\d{2}\/\d{4})", text, flags=re.IGNORECASE)
        if dob_match:
            dob = dob_match.group(1).strip()

        uid_match = re.search(
            r"((?<=\s)\d{4}\s(?<=\s)\d{4}\s(?<=\s)\d{4})(?:\sVID\s*:\s*\d{4}\s\d{4}\s\d{4}\s\d{4})?",
            text,
            flags=re.IGNORECASE,
        )
        if uid_match:
            uid = re.sub(r"\D", "", uid_match.group(1))

        uid = extract_valid_uid_from_text(uid)

        upper_text = text.upper()
        if "FEMALE" in upper_text:
            gender = "Female"
        elif "MALE" in upper_text:
            gender = "Male"

        return ParseResult(name=name, dob=dob, uid=uid, gender=gender)

    def run_with_bytes(self, image_bytes: bytes) -> tuple[ParseResult, str]:
        ocr_result, _ = self.detect_text(image_bytes)
        full_text = self.extract_full_text(ocr_result)
        if not full_text:
            raise ValueError("No text detected by OCR")
        parsed = self.parse_with_regex(full_text)
        if not validate_verhoeff(parsed.uid):
            raise ValueError("No valid UID found using Verhoeff check")
        return parsed, full_text


class SupabaseStore:
    def __init__(self) -> None:
        self.client = self._init_client()
        self.table = SUPABASE_PARSE_TABLE
        self.ready = self._validate_table()

    def _init_client(self):
        if not create_supabase_client:
            return None
        if not SUPABASE_URL or not SUPABASE_KEY:
            return None
        try:
            return create_supabase_client(SUPABASE_URL, SUPABASE_KEY)
        except Exception as exc:
            logger.exception("Supabase client init failed: %s", exc)
            return None

    def _validate_table(self) -> bool:
        if self.client is None:
            return False
        try:
            self.client.table(self.table).select("id").limit(1).execute()
            return True
        except Exception as exc:
            logger.error("Supabase table '%s' is not reachable: %s", self.table, exc)
            return False

    def save_parse(self, parsed: ParseResult, ocr_text: str, source: str, tg_user: TelegramUserCtx | None) -> str:
        if self.client is None:
            return "failed: supabase_not_configured"
        if not self.ready:
            return "failed: parse_table_not_ready"
        if tg_user is None:
            return "failed: user_context_required"

        payload = {
            "telegram_user_id": tg_user.user_id,
            "telegram_username": tg_user.username,
            "name": parsed.name.strip(),
            "dob": parsed.dob.strip(),
            "uid": parsed.uid.strip(),
            "gender": parsed.gender.strip(),
            "ocr_text": ocr_text,
            "source": source,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        missing = [k for k in ("name", "dob", "uid", "gender") if not str(payload.get(k, "")).strip()]
        if missing:
            logger.error("Supabase save skipped: missing required fields: %s", ",".join(missing))
            return f"failed: required_field_missing:{','.join(missing)}"

        try:
            self.client.table(self.table).insert(payload).execute()
            return "saved"
        except Exception as exc:
            logger.exception("Supabase save failed (user=%s uid=%s): %s", tg_user.user_id, payload["uid"], exc)
            return f"failed: {exc}"

    def list_user_parses(self, user_id: int, limit: int, offset: int) -> list[dict[str, Any]]:
        result = (
            self.client.table(self.table)
            .select("id,telegram_user_id,telegram_username,name,dob,uid,gender,source,created_at")
            .eq("telegram_user_id", user_id)
            .order("created_at", desc=True)
            .range(offset, offset + limit - 1)
            .execute()
        )
        return result.data or []

    def admin_search(self, keyword: str, limit: int, offset: int) -> list[dict[str, Any]]:
        base_query = (
            self.client.table(self.table)
            .select("id,telegram_user_id,telegram_username,name,dob,uid,gender,source,created_at")
            .order("created_at", desc=True)
            .range(offset, offset + limit - 1)
        )
        if keyword.strip():
            like_term = f"%{keyword.strip()}%"
            fields = ["telegram_username", "name", "dob", "uid", "gender", "source"]
            or_query = ",".join(f"{f}.ilike.{like_term}" for f in fields)
            result = base_query.or_(or_query).execute()
        else:
            result = base_query.execute()
        return result.data or []


class LocalExcelStore:
    HEADERS = [
        "created_at",
        "telegram_user_id",
        "telegram_username",
        "name",
        "dob",
        "uid",
        "gender",
        "ocr_text",
        "source",
    ]

    def __init__(self, enabled: bool, path: str, sheet: str) -> None:
        self.enabled = enabled
        self.path = path
        self.sheet = sheet
        self.ready = self._init_file() if enabled else False

    def _init_file(self) -> bool:
        try:
            if not self.path:
                return False
            if not os.path.exists(self.path):
                wb = Workbook()
                ws = wb.active
                ws.title = self.sheet
                ws.append(self.HEADERS)
                wb.save(self.path)
                return True

            wb = load_workbook(self.path)
            ws = wb[self.sheet] if self.sheet in wb.sheetnames else wb.active
            first_row = [cell.value for cell in ws[1]]
            return first_row == self.HEADERS
        except Exception as exc:
            logger.exception("Local Excel init failed: %s", exc)
            return False

    def save_parse(self, parsed: ParseResult, ocr_text: str, source: str, tg_user: TelegramUserCtx | None) -> str:
        if not self.enabled:
            return "skipped: disabled"
        if not self.ready:
            return "failed: not_ready"
        if tg_user is None:
            return "failed: user_context_required"

        payload = {
            "created_at": datetime.now(timezone.utc).isoformat(),
            "telegram_user_id": tg_user.user_id,
            "telegram_username": tg_user.username or "",
            "name": parsed.name.strip(),
            "dob": parsed.dob.strip(),
            "uid": parsed.uid.strip(),
            "gender": parsed.gender.strip(),
            "ocr_text": ocr_text,
            "source": source,
        }
        missing = [k for k in ("name", "dob", "uid", "gender") if not payload[k]]
        if missing:
            return f"failed: required_field_missing:{','.join(missing)}"

        try:
            wb = load_workbook(self.path)
            ws = wb[self.sheet] if self.sheet in wb.sheetnames else wb.active
            ws.append([payload.get(col, "") for col in self.HEADERS])
            wb.save(self.path)
            return "saved"
        except Exception as exc:
            logger.exception("Local Excel save failed: %s", exc)
            return f"failed: {exc}"


class StorageManager:
    def __init__(self) -> None:
        self.supabase = SupabaseStore()
        self.local_excel = LocalExcelStore(
            enabled=LOCAL_EXCEL_ENABLED,
            path=LOCAL_EXCEL_FILE,
            sheet=LOCAL_EXCEL_SHEET,
        )

    def save_parse(self, parsed: ParseResult, ocr_text: str, source: str, tg_user: TelegramUserCtx | None) -> dict[str, str]:
        return {
            "supabase": self.supabase.save_parse(parsed, ocr_text, source, tg_user),
            "local_excel": self.local_excel.save_parse(parsed, ocr_text, source, tg_user),
        }

    def list_user_parses(self, user_id: int, limit: int, offset: int) -> list[dict[str, Any]]:
        return self.supabase.list_user_parses(user_id, limit, offset)

    def admin_search(self, keyword: str, limit: int, offset: int) -> list[dict[str, Any]]:
        return self.supabase.admin_search(keyword, limit, offset)

    @property
    def supabase_ready(self) -> bool:
        return bool(self.supabase.client and self.supabase.ready)


parser = OCRParser()
store = StorageManager()


def _startup_preflight() -> None:
    checks_ok: list[str] = []
    checks_warn: list[str] = []
    checks_fail: list[str] = []

    required_env = {
        "TELEGRAM_BOT_TOKEN": TELEGRAM_BOT_TOKEN,
        "TELEGRAM_WEBHOOK_SECRET": TELEGRAM_WEBHOOK_SECRET,
        "SUPABASE_URL": SUPABASE_URL,
        "SUPABASE_KEY_OR_SERVICE_KEY": SUPABASE_KEY,
    }
    for key, val in required_env.items():
        if val:
            checks_ok.append(f"{key}: set")
        else:
            checks_fail.append(f"{key}: missing")

    if MAX_UPLOAD_SIZE_BYTES <= 0:
        checks_fail.append("MAX_UPLOAD_SIZE_BYTES must be > 0")
    else:
        checks_ok.append(f"MAX_UPLOAD_SIZE_BYTES={MAX_UPLOAD_SIZE_BYTES}")

    if not CORS_ALLOW_ORIGINS:
        checks_warn.append("CORS_ALLOW_ORIGINS is empty; default behavior may block browser clients")
    else:
        checks_ok.append(f"CORS_ALLOW_ORIGINS configured ({len(CORS_ALLOW_ORIGINS)} entries)")

    if not ADMIN_TELEGRAM_USER_IDS:
        checks_warn.append("ADMIN_TELEGRAM_USER_IDS is empty; admin search will be inaccessible")
    else:
        checks_ok.append(f"ADMIN_TELEGRAM_USER_IDS loaded ({len(ADMIN_TELEGRAM_USER_IDS)} users)")

    if LOCAL_EXCEL_ENABLED:
        if store.local_excel.ready:
            checks_ok.append(f"Local Excel enabled and ready ({LOCAL_EXCEL_FILE}:{LOCAL_EXCEL_SHEET})")
        else:
            checks_fail.append("Local Excel enabled but initialization failed")
    else:
        checks_ok.append("Local Excel disabled")

    if store.supabase_ready:
        checks_ok.append(f"Supabase table reachable ({SUPABASE_PARSE_TABLE})")
    else:
        checks_fail.append(f"Supabase table not reachable ({SUPABASE_PARSE_TABLE})")

    if TELEGRAM_BOT_TOKEN:
        try:
            bot_info = telegram_api("getMe")
            username = bot_info.get("result", {}).get("username", "unknown")
            checks_ok.append(f"Telegram bot API reachable (@{username})")
        except Exception as exc:
            checks_fail.append(f"Telegram getMe failed: {exc}")

    expected_routes = {
        "/health",
        "/demo",
        "/api/parse",
        "/api/me/parses",
        "/api/admin/search",
        "/telegram/webhook/{webhook_secret}",
    }
    registered_routes = {getattr(r, "path", "") for r in app.router.routes}
    missing_routes = sorted(path for path in expected_routes if path not in registered_routes)
    if missing_routes:
        checks_fail.append(f"Missing required routes: {', '.join(missing_routes)}")
    else:
        checks_ok.append("Required routes registered")

    logger.info("Startup preflight passed checks:\n- %s", "\n- ".join(checks_ok) if checks_ok else "none")
    if checks_warn:
        logger.warning("Startup preflight warnings:\n- %s", "\n- ".join(checks_warn))
    if checks_fail:
        logger.error("Startup preflight failed checks:\n- %s", "\n- ".join(checks_fail))
        raise RuntimeError("Startup preflight failed. Resolve configuration/connectivity issues and restart.")


@asynccontextmanager
async def lifespan(_: FastAPI):
    await run_in_threadpool(_startup_preflight)
    yield


app = FastAPI(
    title="Aadhaar Scanner Unified Service",
    version="2.0.0",
    lifespan=lifespan,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ALLOW_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)


def verify_telegram_init_data(init_data: str) -> dict[str, Any]:
    if not init_data:
        raise HTTPException(status_code=401, detail={"error": "missing_init_data", "message": "Missing Telegram initData."})
    if not TELEGRAM_BOT_TOKEN:
        raise HTTPException(status_code=500, detail={"error": "server_misconfigured", "message": "TELEGRAM_BOT_TOKEN is missing."})

    parsed = urllib.parse.parse_qsl(init_data, keep_blank_values=True)
    data = dict(parsed)
    hash_value = data.pop("hash", None)
    if not hash_value:
        raise HTTPException(status_code=401, detail={"error": "invalid_init_data", "message": "Missing hash in initData."})

    check_string = "\n".join(f"{k}={v}" for k, v in sorted(data.items()))
    secret = hmac.new(b"WebAppData", TELEGRAM_BOT_TOKEN.encode("utf-8"), hashlib.sha256).digest()
    computed = hmac.new(secret, check_string.encode("utf-8"), hashlib.sha256).hexdigest()

    if not hmac.compare_digest(computed, hash_value):
        raise HTTPException(status_code=401, detail={"error": "invalid_init_data", "message": "Invalid Telegram signature."})

    user_raw = data.get("user")
    if not user_raw:
        raise HTTPException(status_code=401, detail={"error": "invalid_init_data", "message": "Missing Telegram user in initData."})

    user = json.loads(user_raw)
    user_id = int(user["id"])
    return {
        "user_id": user_id,
        "username": user.get("username"),
        "first_name": user.get("first_name"),
    }


def auth_user_ctx(authorization: str = Header(default="")) -> TelegramUserCtx:
    if not authorization.lower().startswith("tma "):
        raise HTTPException(status_code=401, detail={"error": "unauthorized", "message": "Expected Authorization: tma <initData>."})
    init_data = authorization[4:].strip()
    info = verify_telegram_init_data(init_data)
    return TelegramUserCtx(
        user_id=info["user_id"],
        username=info.get("username"),
        first_name=info.get("first_name"),
        is_admin=info["user_id"] in ADMIN_TELEGRAM_USER_IDS,
    )


def ensure_admin(user: TelegramUserCtx = Depends(auth_user_ctx)) -> TelegramUserCtx:
    if not user.is_admin:
        raise HTTPException(status_code=403, detail={"error": "forbidden", "message": "Admin access required."})
    return user


async def _parse_image_bytes(image_bytes: bytes) -> tuple[ParseResult, str]:
    try:
        return await run_in_threadpool(parser.run_with_bytes, image_bytes)
    except requests.RequestException:
        raise HTTPException(status_code=502, detail={"error": "ocr_upstream_error", "message": "Azure OCR request failed."})
    except Exception as exc:
        raise HTTPException(status_code=422, detail={"error": "processing_failed", "message": str(exc)}) from exc


@app.middleware("http")
async def security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/parse")
async def parse_image(file: UploadFile = File(...), user: TelegramUserCtx = Depends(auth_user_ctx)) -> JSONResponse:
    if file.content_type not in ALLOWED_CONTENT_TYPES:
        raise HTTPException(status_code=415, detail={"error": "unsupported_media_type", "message": f"Allowed: {sorted(ALLOWED_CONTENT_TYPES)}"})
    ext = _extension(file.filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail={"error": "unsupported_file_extension", "message": f"Allowed: {sorted(ALLOWED_EXTENSIONS)}"})

    image_bytes = await file.read()
    if not image_bytes:
        raise HTTPException(status_code=400, detail={"error": "empty_file", "message": "Uploaded file is empty."})
    if len(image_bytes) > MAX_UPLOAD_SIZE_BYTES:
        raise HTTPException(status_code=413, detail={"error": "file_too_large", "message": f"Max upload is {MAX_UPLOAD_SIZE_BYTES} bytes."})
    if not _looks_like_image(image_bytes):
        raise HTTPException(status_code=415, detail={"error": "invalid_image_signature", "message": "Invalid JPEG/PNG/WEBP payload."})

    parsed, ocr_text = await _parse_image_bytes(image_bytes)
    storage_status = store.save_parse(parsed, ocr_text, source="api", tg_user=user)
    return JSONResponse(
        status_code=200,
        content={
            "status": "success",
            "filename": file.filename,
            "content_type": file.content_type,
            "size_bytes": len(image_bytes),
            "ocr_text": ocr_text,
            "parsed": parsed.model_dump(),
            "storage_status": storage_status,
        },
    )


@app.get("/api/me/parses", response_model=ParseListResponse)
async def my_parses(limit: int = 25, offset: int = 0, user: TelegramUserCtx = Depends(auth_user_ctx)) -> ParseListResponse:
    if limit < 1 or limit > 200:
        raise HTTPException(status_code=400, detail={"error": "invalid_limit", "message": "limit must be between 1 and 200."})
    if offset < 0:
        raise HTTPException(status_code=400, detail={"error": "invalid_offset", "message": "offset must be >= 0."})
    if not store.supabase_ready:
        raise HTTPException(status_code=503, detail={"error": "storage_not_configured", "message": "Supabase is not configured or table is unavailable."})

    records = await run_in_threadpool(store.list_user_parses, user.user_id, limit, offset)
    return ParseListResponse(total=len(records), is_admin=user.is_admin, records=records)


@app.get("/api/admin/search", response_model=ParseListResponse)
async def admin_search(keyword: str = "", limit: int = 50, offset: int = 0, user: TelegramUserCtx = Depends(ensure_admin)) -> ParseListResponse:
    if limit < 1 or limit > 200:
        raise HTTPException(status_code=400, detail={"error": "invalid_limit", "message": "limit must be between 1 and 200."})
    if offset < 0:
        raise HTTPException(status_code=400, detail={"error": "invalid_offset", "message": "offset must be >= 0."})
    if not store.supabase_ready:
        raise HTTPException(status_code=503, detail={"error": "storage_not_configured", "message": "Supabase is not configured or table is unavailable."})

    records = await run_in_threadpool(store.admin_search, keyword, limit, offset)
    return ParseListResponse(total=len(records), is_admin=True, records=records)


async def _process_telegram_message(msg: dict[str, Any]) -> None:
    chat = msg.get("chat") or {}
    chat_id = chat.get("id")
    if chat_id is None:
        return

    sender = msg.get("from") or {}
    user_id = sender.get("id")
    username = sender.get("username")
    first_name = sender.get("first_name") or "User"

    photos = msg.get("photo") or []
    document = msg.get("document") or {}
    file_id = None
    if photos:
        file_id = photos[-1].get("file_id")
    elif document and str(document.get("mime_type", "")).startswith("image/"):
        file_id = document.get("file_id")

    if not file_id:
        send_message(chat_id, f"Hi {_escape_markdown(first_name)}, send me an Aadhaar image.")
        return

    processing = send_message(chat_id, f"Hi {_escape_markdown(first_name)},\n\n📥 Image received\n🧠 Running OCR + Aadhaar parsing...")
    processing_message_id = processing.get("result", {}).get("message_id")

    try:
        image_bytes = get_file_bytes(file_id)
        parsed, ocr_text = await _parse_image_bytes(image_bytes)
        user_ctx = TelegramUserCtx(
            user_id=int(user_id),
            username=username,
            first_name=first_name,
            is_admin=int(user_id) in ADMIN_TELEGRAM_USER_IDS,
        )
        storage_status = store.save_parse(parsed, ocr_text, source="webhook", tg_user=user_ctx)

        supabase_status = storage_status.get("supabase", "unknown")
        excel_status = storage_status.get("local_excel", "unknown")
        compact_storage = "Saved"
        if str(supabase_status).startswith("failed"):
            compact_storage = "Save issue"
        elif str(excel_status).startswith("failed"):
            compact_storage = "Saved (sheet issue)"
        elif "skipped" in str(excel_status):
            compact_storage = "Saved"

        parsed_text = (
            f"*Name:* *{_escape_markdown(parsed.name)}*\n"
            f"*Date of Birth:* {_escape_markdown(parsed.dob)}\n"
            f"*Gender:* {_escape_markdown(parsed.gender)}\n"
            f"*Aadhaar No:* `{_escape_markdown(parsed.uid)}`\n"
            f"*Storage:* {_escape_markdown(compact_storage)}"
        )

        final_text = f"✅ *Aadhaar Parsed Successfully*\n\n{parsed_text}"
        if isinstance(processing_message_id, int):
            edit_message_text(chat_id, processing_message_id, final_text)
        else:
            send_message(chat_id, final_text)
    except Exception as exc:
        logger.exception("Webhook processing failed: %s", exc)
        failure_text = (
            "❌ Could not parse a valid Aadhaar.\n\n"
            "Please send a clearer Aadhaar image with visible number and DOB."
        )
        if isinstance(processing_message_id, int):
            try:
                edit_message_text(chat_id, processing_message_id, failure_text)
            except Exception:
                send_message(chat_id, failure_text)
        else:
            send_message(chat_id, failure_text)


@app.post("/telegram/webhook/{webhook_secret}")
async def telegram_webhook(webhook_secret: str, payload: dict[str, Any]) -> dict[str, bool]:
    if not TELEGRAM_WEBHOOK_SECRET:
        raise HTTPException(status_code=500, detail={"error": "server_misconfigured", "message": "TELEGRAM_WEBHOOK_SECRET is missing."})
    if webhook_secret != TELEGRAM_WEBHOOK_SECRET:
        raise HTTPException(status_code=403, detail={"error": "forbidden", "message": "Invalid webhook secret."})

    message = payload.get("message") or payload.get("edited_message")
    if message:
        await _process_telegram_message(message)
    return {"ok": True}


DEMO_HTML = """<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover" />
  <title>Aadhaar Scanner Mini App</title>
  <script src="https://telegram.org/js/telegram-web-app.js"></script>
  <script src="https://code.iconify.design/iconify-icon/1.0.8/iconify-icon.min.js"></script>
  <style>
    :root {
      --bg: #f4f5f6;
      --surface: #fcfcfb;
      --ink: #212326;
      --muted: #6f747c;
      --line: #e5e7ea;
      --accent: #4b6b8f;
      --accent-soft: #e9eef5;
      --success: #2f7d53;
      --danger: #a44747;
      --radius: 14px;
      --nav-h: 56px;
    }
    [data-theme="warm"] {
      --bg: #f5f3ef;
      --surface: #fffdf9;
      --ink: #24211f;
      --muted: #786f67;
      --line: #ebe3d8;
      --accent: #7d6b5a;
      --accent-soft: #f2ece4;
    }
    [data-theme="soft-dark"] {
      --bg: #1f2124;
      --surface: #2a2d31;
      --ink: #f5f4f2;
      --muted: #b8bcc3;
      --line: #3a3f45;
      --accent: #8aa4c2;
      --accent-soft: #313a44;
      --success: #80c89f;
      --danger: #e29595;
    }

    * { box-sizing: border-box; }
    html, body { height: 100%; }
    body {
      margin: 0;
      background: radial-gradient(circle at top left, var(--accent-soft), var(--bg) 42%);
      color: var(--ink);
      font-family: "Inter", "Segoe UI", ui-sans-serif, system-ui, sans-serif;
      -webkit-font-smoothing: antialiased;
      text-rendering: optimizeLegibility;
    }

    .app {
      max-width: 1080px;
      margin: 0 auto;
      min-height: 100%;
      display: grid;
      grid-template-rows: auto 1fr auto;
      padding-bottom: calc(var(--nav-h) + 12px);
    }

    .topbar {
      position: sticky;
      top: 0;
      z-index: 20;
      backdrop-filter: blur(12px);
      background: color-mix(in srgb, var(--surface) 88%, transparent);
      border-bottom: 1px solid var(--line);
      padding: env(safe-area-inset-top) 12px 10px;
    }

    .topbar-row {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 10px;
      min-height: 42px;
    }

    .brand {
      display: inline-flex;
      align-items: center;
      gap: 8px;
      font-weight: 700;
      letter-spacing: -.01em;
      font-size: 1rem;
    }

    .icon-btn, .btn {
      border-radius: 11px;
      border: 1px solid var(--line);
      background: var(--surface);
      color: var(--ink);
      padding: 9px 12px;
      font-size: .9rem;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      gap: 6px;
      white-space: nowrap;
    }

    .btn.primary {
      background: var(--accent);
      color: #fff;
      border-color: var(--accent);
      font-weight: 600;
    }

    .btn:disabled, .icon-btn:disabled { opacity: .6; cursor: not-allowed; }

    .whoami {
      color: var(--muted);
      font-size: .84rem;
      margin-top: 6px;
      overflow: hidden;
      text-overflow: ellipsis;
      white-space: nowrap;
    }

    .popup {
      position: fixed;
      inset: 0;
      z-index: 40;
      display: none;
      align-items: flex-start;
      justify-content: flex-end;
      padding: calc(env(safe-area-inset-top) + 56px) 12px 12px;
      background: rgba(0,0,0,.18);
    }

    .popup.open { display: flex; }

    .popup-menu {
      width: min(240px, 88vw);
      background: var(--surface);
      border: 1px solid var(--line);
      border-radius: 12px;
      box-shadow: 0 14px 30px rgba(0,0,0,.12);
      padding: 8px;
      display: grid;
      gap: 6px;
    }

    .menu-item {
      width: 100%;
      text-align: left;
      border: 1px solid var(--line);
      border-radius: 10px;
      background: var(--surface);
      color: var(--ink);
      padding: 10px 11px;
      font-size: .88rem;
      display: inline-flex;
      align-items: center;
      gap: 8px;
      cursor: pointer;
    }

    .content {
      padding: 12px;
      display: grid;
      gap: 12px;
      align-content: start;
    }

    .snap-row {
      display: grid;
      grid-auto-flow: column;
      grid-auto-columns: minmax(84vw, 1fr);
      gap: 12px;
      overflow-x: auto;
      scroll-snap-type: x mandatory;
      padding-bottom: 4px;
      -webkit-overflow-scrolling: touch;
    }

    .snap-row > .card {
      scroll-snap-align: start;
    }

    .card {
      background: var(--surface);
      border: 1px solid var(--line);
      border-radius: var(--radius);
      box-shadow: 0 8px 22px rgba(0,0,0,.04);
      padding: 14px;
    }

    .card-head {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 10px;
      margin-bottom: 10px;
    }

    .title {
      margin: 0;
      font-size: .96rem;
      font-weight: 700;
      display: inline-flex;
      align-items: center;
      gap: 7px;
      min-width: 0;
    }

    .badge {
      border: 1px solid var(--line);
      background: color-mix(in srgb, var(--surface) 88%, var(--accent-soft));
      border-radius: 999px;
      padding: 4px 9px;
      font-size: .75rem;
      color: var(--muted);
      flex-shrink: 0;
    }

    .controls {
      display: grid;
      grid-template-columns: 1fr;
      gap: 9px;
    }

    .control-row {
      display: flex;
      gap: 8px;
      align-items: center;
      flex-wrap: nowrap;
    }

    .input {
      width: 100%;
      min-width: 0;
      border: 1px solid var(--line);
      background: var(--surface);
      color: var(--ink);
      border-radius: 11px;
      padding: 10px 12px;
      font-size: .9rem;
    }

    .status { font-size: .88rem; font-weight: 600; margin-top: 8px; color: var(--muted); }
    .ok { color: var(--success); }
    .bad { color: var(--danger); }

    .collapsible {
      border: 1px solid var(--line);
      border-radius: 12px;
      overflow: hidden;
      margin-top: 10px;
    }

    .collapse-head {
      width: 100%;
      background: color-mix(in srgb, var(--surface) 92%, var(--accent-soft));
      border: 0;
      color: var(--ink);
      padding: 10px 12px;
      font-weight: 600;
      text-align: left;
      display: flex;
      align-items: center;
      justify-content: space-between;
      cursor: pointer;
    }

    .collapse-body {
      max-height: 0;
      overflow: hidden;
      transition: max-height .22s ease;
      background: var(--surface);
    }

    .collapsible.open .collapse-body { max-height: 280px; }

    pre {
      margin: 0;
      padding: 12px;
      max-height: 260px;
      overflow: auto;
      font-size: .78rem;
      background: color-mix(in srgb, var(--surface) 94%, var(--accent-soft));
    }

    .table-wrap {
      overflow: auto;
      border: 1px solid var(--line);
      border-radius: 12px;
      max-height: 52vh;
    }

    table {
      width: 100%;
      border-collapse: separate;
      border-spacing: 0;
      min-width: 660px;
      font-size: .84rem;
    }

    th, td {
      padding: 9px;
      border-bottom: 1px solid var(--line);
      text-align: left;
      white-space: nowrap;
      vertical-align: middle;
    }

    th {
      position: sticky;
      top: 0;
      z-index: 1;
      background: color-mix(in srgb, var(--surface) 92%, var(--accent-soft));
      color: var(--muted);
      font-weight: 600;
    }

    tr:last-child td { border-bottom: 0; }

    .pager {
      margin-top: 10px;
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 8px;
    }

    .pager .group { display: inline-flex; gap: 8px; align-items: center; }

    .bottom-nav {
      position: fixed;
      left: 0;
      right: 0;
      bottom: 0;
      z-index: 30;
      height: var(--nav-h);
      background: color-mix(in srgb, var(--surface) 88%, transparent);
      backdrop-filter: blur(10px);
      border-top: 1px solid var(--line);
      display: flex;
      justify-content: center;
      padding-bottom: env(safe-area-inset-bottom);
    }

    .bottom-wrap {
      width: min(1080px, 100%);
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 8px;
      padding: 8px 12px;
    }

    .nav-btn {
      border: 1px solid var(--line);
      background: var(--surface);
      color: var(--muted);
      border-radius: 10px;
      font-size: .8rem;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      gap: 6px;
    }

    .nav-btn.active {
      color: var(--ink);
      border-color: var(--accent);
      background: color-mix(in srgb, var(--surface) 74%, var(--accent-soft));
    }

    @media (min-width: 860px) {
      .content { padding: 16px; }
      .snap-row {
        grid-auto-flow: row;
        grid-template-columns: 1fr 1fr;
        overflow: visible;
        scroll-snap-type: none;
      }
      .snap-row > .card { min-height: 100%; }
      .bottom-nav { display: none; }
      .app { padding-bottom: 18px; }
      .controls { grid-template-columns: 1fr auto; align-items: center; }
      .controls .full { grid-column: 1 / -1; }
    }
  </style>
</head>
<body data-theme="default">
  <div class="app">
    <header class="topbar">
      <div class="topbar-row">
        <div class="brand"><iconify-icon icon="solar:shield-user-outline"></iconify-icon> Aadhaar Scanner</div>
        <button id="menuBtn" class="icon-btn" aria-label="Menu"><iconify-icon icon="solar:hamburger-menu-outline"></iconify-icon></button>
      </div>
      <div class="whoami" id="whoami">Authenticating...</div>
    </header>

    <main class="content">
      <div class="snap-row">
        <section class="card" id="card-upload">
          <div class="card-head">
            <h2 class="title"><iconify-icon icon="solar:document-add-outline"></iconify-icon> Parse Aadhaar</h2>
            <span class="badge">Secure</span>
          </div>
          <div class="controls">
            <div class="control-row full">
              <input id="file" class="input" type="file" accept=".jpg,.jpeg,.png,.webp,image/jpeg,image/png,image/webp" />
            </div>
            <div class="control-row">
              <button id="parseBtn" class="btn primary"><iconify-icon icon="solar:play-circle-outline"></iconify-icon> Parse</button>
            </div>
          </div>
          <div id="parseStatus" class="status">Idle</div>
          <div class="collapsible" id="parseCollapse">
            <button class="collapse-head" id="parseToggle">
              <span>Parsed JSON</span>
              <iconify-icon icon="solar:alt-arrow-down-outline"></iconify-icon>
            </button>
            <div class="collapse-body">
              <pre id="parseResult">{}</pre>
            </div>
          </div>
        </section>

        <section class="card" id="card-mine">
          <div class="card-head">
            <h2 class="title"><iconify-icon icon="solar:history-outline"></iconify-icon> My Records</h2>
            <span class="badge" id="myCount">0 records</span>
          </div>
          <div class="table-wrap">
            <table>
              <thead>
                <tr><th>Created At</th><th>Name</th><th>DOB</th><th>UID</th><th>Gender</th><th>Source</th></tr>
              </thead>
              <tbody id="myRows"></tbody>
            </table>
          </div>
          <div class="pager">
            <span class="muted" id="myPageInfo">Page 1</span>
            <div class="group">
              <button id="myPrev" class="btn">Prev</button>
              <button id="myNext" class="btn">Next</button>
            </div>
          </div>
        </section>
      </div>

      <section class="card" id="adminCard" style="display:none">
        <div class="card-head">
          <h2 class="title"><iconify-icon icon="solar:magnifer-outline"></iconify-icon> Admin Search</h2>
          <span class="badge" id="adminCount">0 results</span>
        </div>
        <div class="controls">
          <div class="control-row full">
            <input id="keyword" class="input" placeholder="Search username, name, UID, gender" />
          </div>
          <div class="control-row">
            <button id="searchBtn" class="btn primary"><iconify-icon icon="solar:magnifer-outline"></iconify-icon> Search</button>
            <button id="clearSearchBtn" class="btn"><iconify-icon icon="solar:close-circle-outline"></iconify-icon> Clear</button>
          </div>
        </div>
        <div class="table-wrap" style="margin-top:10px">
          <table>
            <thead>
              <tr><th>Created At</th><th>User ID</th><th>Username</th><th>Name</th><th>DOB</th><th>UID</th><th>Gender</th></tr>
            </thead>
            <tbody id="adminRows"></tbody>
          </table>
        </div>
        <div class="pager">
          <span class="muted" id="adminPageInfo">Page 1</span>
          <div class="group">
            <button id="adminPrev" class="btn">Prev</button>
            <button id="adminNext" class="btn">Next</button>
          </div>
        </div>
      </section>
    </main>

    <nav class="bottom-nav">
      <div class="bottom-wrap">
        <button class="nav-btn active" data-target="card-upload"><iconify-icon icon="solar:document-add-outline"></iconify-icon> Parse</button>
        <button class="nav-btn" data-target="card-mine"><iconify-icon icon="solar:history-outline"></iconify-icon> Records</button>
        <button class="nav-btn" data-target="adminCard"><iconify-icon icon="solar:magnifer-outline"></iconify-icon> Admin</button>
      </div>
    </nav>
  </div>

  <div id="popup" class="popup" aria-hidden="true">
    <div class="popup-menu" role="menu">
      <button class="menu-item" data-theme-btn="default"><iconify-icon icon="solar:palette-round-outline"></iconify-icon> Theme: Default</button>
      <button class="menu-item" data-theme-btn="warm"><iconify-icon icon="solar:palette-round-outline"></iconify-icon> Theme: Warm</button>
      <button class="menu-item" data-theme-btn="soft-dark"><iconify-icon icon="solar:palette-round-outline"></iconify-icon> Theme: Soft Dark</button>
      <button class="menu-item" id="refreshAll"><iconify-icon icon="solar:refresh-outline"></iconify-icon> Refresh Data</button>
    </div>
  </div>

<script>
const tg = window.Telegram?.WebApp;
if (tg) tg.ready();
const initData = tg?.initData || "";
const authHeader = { "Authorization": `tma ${initData}` };

const parseBtn = document.getElementById('parseBtn');
const fileInput = document.getElementById('file');
const parseStatus = document.getElementById('parseStatus');
const parseResult = document.getElementById('parseResult');
const whoami = document.getElementById('whoami');

const myLimit = 20;
let myOffset = 0;
const adminLimit = 20;
let adminOffset = 0;
let adminKeyword = '';

function formatTimestamp(v, withTime = true) {
  if (!v) return '';
  const d = new Date(v);
  if (Number.isNaN(d.getTime())) return v;
  const dd = String(d.getDate()).padStart(2, '0');
  const mm = String(d.getMonth() + 1).padStart(2, '0');
  const yyyy = d.getFullYear();
  if (!withTime) return `${dd}-${mm}-${yyyy}`;
  const hh = String(d.getHours()).padStart(2, '0');
  const min = String(d.getMinutes()).padStart(2, '0');
  const ss = String(d.getSeconds()).padStart(2, '0');
  return `${dd}-${mm}-${yyyy} ${hh}:${min}:${ss}`;
}

function setStatus(el, text, ok = true) {
  el.textContent = text;
  el.classList.remove('ok', 'bad');
  el.classList.add(ok ? 'ok' : 'bad');
}

function setPageInfo(elId, offset, limit) {
  const page = Math.floor(offset / limit) + 1;
  document.getElementById(elId).textContent = `Page ${page}`;
}

function myRowHtml(r) {
  return `<tr>
    <td>${formatTimestamp(r.created_at, true)}</td>
    <td>${r.name || ''}</td>
    <td>${r.dob || ''}</td>
    <td>${r.uid || ''}</td>
    <td>${r.gender || ''}</td>
    <td>${r.source || ''}</td>
  </tr>`;
}

function adminRowHtml(r) {
  return `<tr>
    <td>${formatTimestamp(r.created_at, true)}</td>
    <td>${r.telegram_user_id || ''}</td>
    <td>${r.telegram_username || ''}</td>
    <td>${r.name || ''}</td>
    <td>${r.dob || ''}</td>
    <td>${r.uid || ''}</td>
    <td>${r.gender || ''}</td>
  </tr>`;
}

async function loadMine() {
  const res = await fetch(`/api/me/parses?limit=${myLimit}&offset=${myOffset}`, { headers: authHeader });
  const data = await res.json();
  if (!res.ok) throw new Error(data?.detail?.message || 'Failed to load records');
  const rows = data.records || [];
  document.getElementById('myRows').innerHTML = rows.map(myRowHtml).join('') || '<tr><td colspan="6">No records</td></tr>';
  document.getElementById('myCount').textContent = `${rows.length} records`;
  document.getElementById('myPrev').disabled = myOffset === 0;
  document.getElementById('myNext').disabled = rows.length < myLimit;
  setPageInfo('myPageInfo', myOffset, myLimit);

  whoami.textContent = `Authenticated. Admin: ${data.is_admin ? 'Yes' : 'No'}`;
  const adminCard = document.getElementById('adminCard');
  adminCard.style.display = data.is_admin ? 'block' : 'none';

  const adminNavBtn = document.querySelector('.nav-btn[data-target="adminCard"]');
  if (adminNavBtn) adminNavBtn.style.display = data.is_admin ? 'inline-flex' : 'none';

  if (data.is_admin) await loadAdmin();
}

async function loadAdmin() {
  const url = `/api/admin/search?keyword=${encodeURIComponent(adminKeyword)}&limit=${adminLimit}&offset=${adminOffset}`;
  const res = await fetch(url, { headers: authHeader });
  const data = await res.json();
  const tbody = document.getElementById('adminRows');
  if (!res.ok) {
    tbody.innerHTML = `<tr><td colspan="7">${data?.detail?.message || 'Search failed'}</td></tr>`;
    document.getElementById('adminCount').textContent = '0 results';
    return;
  }
  const rows = data.records || [];
  tbody.innerHTML = rows.map(adminRowHtml).join('') || '<tr><td colspan="7">No results</td></tr>';
  document.getElementById('adminCount').textContent = `${rows.length} results`;
  document.getElementById('adminPrev').disabled = adminOffset === 0;
  document.getElementById('adminNext').disabled = rows.length < adminLimit;
  setPageInfo('adminPageInfo', adminOffset, adminLimit);
}

async function parseNow() {
  const f = fileInput.files?.[0];
  if (!f) {
    setStatus(parseStatus, 'Choose a file first', false);
    return;
  }
  parseBtn.disabled = true;
  setStatus(parseStatus, 'Processing...', true);
  const fd = new FormData();
  fd.append('file', f);
  try {
    const res = await fetch('/api/parse', { method: 'POST', headers: authHeader, body: fd });
    const data = await res.json();
    parseResult.textContent = JSON.stringify(data, null, 2);
    if (!res.ok) throw new Error(data?.detail?.message || 'Parse failed');
    setStatus(parseStatus, 'Parsed successfully', true);
    myOffset = 0;
    await loadMine();
  } catch (e) {
    setStatus(parseStatus, e.message, false);
  } finally {
    parseBtn.disabled = false;
  }
}

function togglePopup(force) {
  const popup = document.getElementById('popup');
  const open = force !== undefined ? force : !popup.classList.contains('open');
  popup.classList.toggle('open', open);
  popup.setAttribute('aria-hidden', String(!open));
}

function setupInteractions() {
  document.getElementById('parseToggle').addEventListener('click', () => {
    document.getElementById('parseCollapse').classList.toggle('open');
  });

  document.getElementById('myPrev').addEventListener('click', async () => {
    myOffset = Math.max(0, myOffset - myLimit);
    await loadMine();
  });
  document.getElementById('myNext').addEventListener('click', async () => {
    myOffset += myLimit;
    await loadMine();
  });
  document.getElementById('adminPrev').addEventListener('click', async () => {
    adminOffset = Math.max(0, adminOffset - adminLimit);
    await loadAdmin();
  });
  document.getElementById('adminNext').addEventListener('click', async () => {
    adminOffset += adminLimit;
    await loadAdmin();
  });

  document.getElementById('searchBtn').addEventListener('click', async () => {
    adminKeyword = (document.getElementById('keyword').value || '').trim();
    adminOffset = 0;
    await loadAdmin();
  });
  document.getElementById('clearSearchBtn').addEventListener('click', async () => {
    document.getElementById('keyword').value = '';
    adminKeyword = '';
    adminOffset = 0;
    await loadAdmin();
  });

  parseBtn.addEventListener('click', parseNow);

  document.querySelectorAll('[data-theme-btn]').forEach(btn => {
    btn.addEventListener('click', () => {
      document.body.setAttribute('data-theme', btn.getAttribute('data-theme-btn'));
      togglePopup(false);
    });
  });

  const menuBtn = document.getElementById('menuBtn');
  const popup = document.getElementById('popup');
  menuBtn.addEventListener('click', () => togglePopup());
  popup.addEventListener('click', (e) => {
    if (e.target === popup) togglePopup(false);
  });

  document.getElementById('refreshAll').addEventListener('click', async () => {
    togglePopup(false);
    await loadMine();
  });

  document.querySelectorAll('.nav-btn[data-target]').forEach(btn => {
    btn.addEventListener('click', () => {
      document.querySelectorAll('.nav-btn').forEach(x => x.classList.remove('active'));
      btn.classList.add('active');
      const target = document.getElementById(btn.getAttribute('data-target'));
      if (target) target.scrollIntoView({ behavior: 'smooth', block: 'start' });
    });
  });
}

setupInteractions();
loadMine().catch(err => {
  whoami.textContent = `Auth failed: ${err.message}`;
  whoami.classList.add('bad');
});
</script>
</body>
</html>"""


@app.get("/demo", response_class=HTMLResponse)
async def demo_page() -> HTMLResponse:
    return HTMLResponse(content=DEMO_HTML)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="0.0.0.0", port=int(os.getenv("PORT", "8000")), reload=False)
